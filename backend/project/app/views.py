import os
import io
import time
import zipfile
from pathlib import Path
from uuid import uuid4

from django.conf import settings
import fitz  # PyMuPDF
from PIL import Image
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Pt
from pptx import Presentation
from concurrent.futures import ThreadPoolExecutor
from rest_framework.decorators import api_view
from rest_framework.response import Response

# Ensure media folder exists
os.makedirs(settings.MEDIA_ROOT, exist_ok=True)

_last_purge_time = 0


def _purge_old_media_files(max_age_hours=1):
    global _last_purge_time
    now = time.time()
    if now - _last_purge_time < 900:  # Purge max once every 15 mins
        return
    _last_purge_time = now
    try:
        cutoff = now - (max_age_hours * 3600)
        media_path = Path(settings.MEDIA_ROOT)
        for item in media_path.iterdir():
            if item.is_file() and item.stat().st_mtime < cutoff:
                try:
                    item.unlink(missing_ok=True)
                except Exception:
                    pass
    except Exception:
        pass


def _build_unique_paths(upload_name, output_extension):
    _purge_old_media_files()
    original_name = Path(upload_name).name
    stem = Path(original_name).stem or "file"
    suffix = Path(original_name).suffix.lower()
    unique_id = uuid4().hex[:8]

    input_name = f"{stem}-{unique_id}-in{suffix}"
    output_name = f"{stem}-{unique_id}-out{output_extension}"

    return (
        Path(settings.MEDIA_ROOT) / input_name,
        Path(settings.MEDIA_ROOT) / output_name,
    )


def _save_uploaded_file(uploaded_file, destination):
    if hasattr(uploaded_file, 'seek'):
        try:
            uploaded_file.seek(0)
        except Exception:
            pass
    with open(destination, "wb+") as target:
        for chunk in uploaded_file.chunks():
            target.write(chunk)


def _cleanup_files(*paths):
    for path in paths:
        if path and Path(path).exists():
            try:
                Path(path).unlink(missing_ok=True)
            except Exception:
                pass


def _get_single_input_path(request, param_name='file'):
    """
    Returns (input_path, original_filename, is_temp_file).
    Supports:
    1. Pre-assembled file path sent in request.data / request.POST ('file_path' or param_name + '_path')
    2. Direct file upload in request.FILES.get(param_name) or request.FILES.get('file')
    """
    _purge_old_media_files()
    file_path_str = (
        request.data.get('file_path') or 
        request.POST.get('file_path') or 
        request.data.get(f"{param_name}_path") or 
        request.POST.get(f"{param_name}_path")
    )
    if file_path_str:
        p = Path(file_path_str)
        if p.exists():
            return p, p.name, False

    uploaded = request.FILES.get(param_name) or request.FILES.get('file')
    if uploaded:
        original_name = Path(uploaded.name).name
        stem = Path(original_name).stem or "file"
        suffix = Path(original_name).suffix.lower()
        unique_id = uuid4().hex[:8]
        target_path = Path(settings.MEDIA_ROOT) / f"{stem}-{unique_id}-in{suffix}"
        _save_uploaded_file(uploaded, target_path)
        return target_path, original_name, True

    return None, None, False


def _get_multiple_input_paths(request, param_name='files'):
    """
    Returns list of tuples: [(input_path, original_filename, is_temp_file), ...]
    Supports:
    1. Pre-assembled file paths sent in request.data / POST ('file_paths' or list)
    2. Direct file uploads in request.FILES.getlist(param_name) or request.FILES.getlist('files')
    """
    _purge_old_media_files()
    results = []

    file_paths = (
        request.data.get('file_paths') or 
        request.POST.get('file_paths')
    )
    if isinstance(file_paths, str):
        import json
        try:
            file_paths = json.loads(file_paths)
        except Exception:
            file_paths = [file_paths]

    if isinstance(file_paths, list) and len(file_paths) > 0:
        for fp in file_paths:
            if fp and isinstance(fp, str):
                p = Path(fp)
                if p.exists():
                    results.append((p, p.name, False))
        if results:
            return results

    files = request.FILES.getlist(param_name) or request.FILES.getlist('files') or ([request.FILES.get('file')] if request.FILES.get('file') else [])
    for uploaded in files:
        if not uploaded:
            continue
        original_name = Path(uploaded.name).name
        stem = Path(original_name).stem or "file"
        suffix = Path(original_name).suffix.lower()
        unique_id = uuid4().hex[:8]
        target_path = Path(settings.MEDIA_ROOT) / f"{stem}-{unique_id}-in{suffix}"
        _save_uploaded_file(uploaded, target_path)
        results.append((target_path, original_name, True))

    return results


def _pdf_has_extractable_text(pdf_path):
    with fitz.open(pdf_path) as pdf_document:
        for page in pdf_document:
            if page.get_text("text").strip():
                return True
    return False


def _pdf_to_docx_pure_python(pdf_path, docx_path):
    pdf_doc = fitz.open(pdf_path)
    document = Document()

    try:
        if pdf_doc.page_count == 0:
            raise ValueError("The uploaded PDF document has no pages.")

        first_section = document.sections[0]
        first_section.top_margin = Pt(36)
        first_section.bottom_margin = Pt(36)
        first_section.left_margin = Pt(36)
        first_section.right_margin = Pt(36)

        total_extracted_chars = 0

        for page_idx, page in enumerate(pdf_doc):
            if page_idx > 0:
                sec = document.add_section(WD_SECTION.NEW_PAGE)
                sec.top_margin = Pt(36)
                sec.bottom_margin = Pt(36)
                sec.left_margin = Pt(36)
                sec.right_margin = Pt(36)

            text_blocks = page.get_text("blocks")
            page_char_count = 0

            if text_blocks:
                for b in text_blocks:
                    if len(b) >= 7 and b[6] == 0:
                        block_text = b[4].strip()
                        if block_text:
                            page_char_count += len(block_text)
                            for line in block_text.split("\n"):
                                line_clean = line.strip()
                                if line_clean:
                                    p = document.add_paragraph()
                                    run = p.add_run(line_clean)
                                    run.font.name = "Calibri"
                                    if len(line_clean) < 45 and (line_clean.isupper() or not line_clean.endswith('.')):
                                        run.font.size = Pt(13)
                                        run.font.bold = True
                                    else:
                                        run.font.size = Pt(11)

            total_extracted_chars += page_char_count

            if page_char_count == 0:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                img_bytes = pixmap.tobytes("png")
                sec = document.sections[-1]
                usable_w = sec.page_width - sec.left_margin - sec.right_margin
                document.add_picture(io.BytesIO(img_bytes), width=usable_w)
                del pixmap, img_bytes

        if total_extracted_chars < 10:
            document = Document()
            sec = document.sections[0]
            sec.top_margin = Pt(18)
            sec.bottom_margin = Pt(18)
            sec.left_margin = Pt(18)
            sec.right_margin = Pt(18)

            for idx, page in enumerate(pdf_doc):
                if idx > 0:
                    sec = document.add_section(WD_SECTION.NEW_PAGE)
                    sec.top_margin = Pt(18)
                    sec.bottom_margin = Pt(18)
                    sec.left_margin = Pt(18)
                    sec.right_margin = Pt(18)
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                img_bytes = pixmap.tobytes("png")
                usable_w = sec.page_width - sec.left_margin - sec.right_margin
                document.add_picture(io.BytesIO(img_bytes), width=usable_w)
                del pixmap, img_bytes

        document.save(str(docx_path))
    finally:
        pdf_doc.close()


def _docx_to_pdf_pure_python(docx_path, pdf_path):
    pdf_doc = fitz.open()
    page_w, page_h = 595.0, 842.0
    margin = 54.0
    page = pdf_doc.new_page(width=page_w, height=page_h)
    y_cursor = margin

    try:
        doc = Document(docx_path)
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                y_cursor += 10.0
                continue
                
            font_size = 11.0
            line_height = 15.0
            if hasattr(p, 'style') and p.style and p.style.name.startswith('Heading 1'):
                font_size = 18.0
                line_height = 22.0
            elif hasattr(p, 'style') and p.style and p.style.name.startswith('Heading 2'):
                font_size = 14.0
                line_height = 18.0

            if y_cursor + line_height > page_h - margin:
                page = pdf_doc.new_page(width=page_w, height=page_h)
                y_cursor = margin

            page.insert_text(fitz.Point(margin, y_cursor), text[:120], fontsize=font_size, color=(0.1, 0.1, 0.1))
            y_cursor += line_height
    except Exception:
        try:
            with zipfile.ZipFile(docx_path, 'r') as z:
                xml_content = z.read("word/document.xml").decode("utf-8", errors="ignore")
                import re
                texts = re.findall(r'<w:t[^>]*>(.*?)</w:t>', xml_content)
                if texts:
                    full_txt = " ".join(texts)
                    page.insert_text(fitz.Point(margin, y_cursor), full_txt[:500], fontsize=11.0)
        except Exception:
            pass

    if pdf_doc.page_count == 0:
        pdf_doc.new_page(width=page_w, height=page_h)

    pdf_doc.save(str(pdf_path), deflate=True)
    pdf_doc.close()


def _pptx_to_pdf_pure_python(pptx_path, pdf_path):
    pdf_doc = fitz.open()
    slide_w, slide_h = 960.0, 540.0
    page = pdf_doc.new_page(width=slide_w, height=slide_h)
    y_cursor = 40.0

    try:
        prs = Presentation(pptx_path)
        for slide in prs.slides:
            if y_cursor > slide_h - 40:
                page = pdf_doc.new_page(width=slide_w, height=slide_h)
                y_cursor = 40.0

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            page.insert_text(fitz.Point(50, y_cursor), text[:120], fontsize=13.0, color=(0.1, 0.1, 0.2))
                            y_cursor += 18.0
    except Exception:
        try:
            with zipfile.ZipFile(pptx_path, 'r') as z:
                slide_files = [f for f in z.namelist() if f.startswith("ppt/slides/slide")]
                import re
                for s_file in slide_files:
                    xml_content = z.read(s_file).decode("utf-8", errors="ignore")
                    texts = re.findall(r'<a:t[^>]*>(.*?)</a:t>', xml_content)
                    if texts:
                        page.insert_text(fitz.Point(50, y_cursor), " ".join(texts)[:300], fontsize=13.0)
                        y_cursor += 25.0
        except Exception:
            pass

    if pdf_doc.page_count == 0:
        pdf_doc.new_page(width=slide_w, height=slide_h)

    pdf_doc.save(str(pdf_path), deflate=True)
    pdf_doc.close()


def _compress_zip_media(input_path, output_path, media_prefix="media/", extreme=True):
    if not zipfile.is_zipfile(input_path):
        import shutil
        shutil.copy2(input_path, output_path)
        return

    quality = 30 if extreme else 55
    max_dim = 900 if extreme else 1450

    with zipfile.ZipFile(input_path, 'r') as in_zip:
        with zipfile.ZipFile(output_path, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=5) as out_zip:
            for item in in_zip.infolist():
                data = in_zip.read(item.filename)
                
                if Path(item.filename).suffix.lower() in ['.jpg', '.jpeg', '.png', '.bmp']:
                    try:
                        img = Image.open(io.BytesIO(data))
                        img_format = img.format or ("JPEG" if Path(item.filename).suffix.lower() in ['.jpg', '.jpeg'] else "PNG")
                        
                        img.thumbnail((max_dim, max_dim), Image.Resampling.BILINEAR)

                        out_buffer = io.BytesIO()
                        
                        if img.mode in ("RGBA", "P") or extreme:
                            img = img.convert("RGB")
                            img_format = "JPEG"
                            
                        if img_format.upper() in ["JPEG", "JPG"]:
                            img.save(out_buffer, format="JPEG", quality=quality)
                        elif img_format.upper() == "PNG":
                            if extreme or img.mode != "RGB":
                                img = img.convert("RGB")
                                img.save(out_buffer, format="JPEG", quality=quality)
                            else:
                                img.save(out_buffer, format="PNG", compress_level=4)
                        else:
                            img.save(out_buffer, format=img_format)
                            
                        compressed_data = out_buffer.getvalue()
                        if len(compressed_data) < len(data):
                            data = compressed_data
                        img.close()
                    except Exception:
                        pass
                        
                out_zip.writestr(item, data)


# ---------------- HEALTH CHECK ----------------
@api_view(['GET'])
def health_check(request):
    _purge_old_media_files()
    return Response({
        "status": "ok",
        "service": "OmniConvert Engine API",
        "active_tools": 15,
        "server_time": time.strftime("%Y-%m-%d %H:%M:%S"),
    })


# ---------------- UPLOAD CHUNK ----------------
@api_view(['POST'])
def upload_chunk(request):
    upload_id = request.POST.get('upload_id') or request.data.get('upload_id')
    chunk_index = request.POST.get('chunk_index') or request.data.get('chunk_index')
    total_chunks = request.POST.get('total_chunks') or request.data.get('total_chunks')
    filename = request.POST.get('filename') or request.data.get('filename') or "file.bin"
    chunk_file = request.FILES.get('chunk')

    if not upload_id or chunk_index is None or not chunk_file:
        return Response({"error": "Missing upload parameters or chunk file"}, status=400)

    try:
        chunk_index = int(chunk_index)
        total_chunks = int(total_chunks) if total_chunks else 1
    except (ValueError, TypeError):
        return Response({"error": "Invalid chunk index or total chunks"}, status=400)

    safe_filename = Path(filename).name
    stem = Path(safe_filename).stem or "upload"
    suffix = Path(safe_filename).suffix.lower()

    assembled_name = f"chunked-{upload_id[:12]}-{stem}{suffix}"
    assembled_path = Path(settings.MEDIA_ROOT) / assembled_name

    try:
        _purge_old_media_files()
        mode = "wb" if chunk_index == 0 else "ab"

        with open(assembled_path, mode) as target:
            for chunk_data in chunk_file.chunks():
                target.write(chunk_data)

        if chunk_index == total_chunks - 1:
            return Response({
                "status": "complete",
                "file_path": str(assembled_path),
                "filename": safe_filename,
                "file_size": assembled_path.stat().st_size
            })
        else:
            return Response({
                "status": "chunk_received",
                "chunk_index": chunk_index,
                "total_chunks": total_chunks
            })
    except Exception as exc:
        return Response({"error": f"Chunk upload failed: {str(exc)}"}, status=500)


# ---------------- 1. PDF → WORD ----------------
@api_view(['POST'])
def pdf_to_word(request):
    pdf_path, original_name, is_temp = _get_single_input_path(request, 'file')

    if not pdf_path or not pdf_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if pdf_path.suffix.lower() != ".pdf":
        if is_temp: _cleanup_files(pdf_path)
        return Response({"error": "Please upload a PDF file"}, status=400)

    _, docx_path = _build_unique_paths(original_name, ".docx")

    try:
        _pdf_to_docx_pure_python(pdf_path, docx_path)
        message = "PDF converted to Word document (.docx) successfully"
    except Exception as exc:
        _cleanup_files(docx_path)
        return Response(
            {"error": f"PDF to Word conversion failed: {str(exc)}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(pdf_path)

    return Response({
        "message": message,
        "file": settings.MEDIA_URL + docx_path.name
    })


# ---------------- 2. WORD → PDF ----------------
@api_view(['POST'])
def word_to_pdf(request):
    word_path, original_name, is_temp = _get_single_input_path(request, 'file')

    if not word_path or not word_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if word_path.suffix.lower() not in [".docx", ".doc"]:
        if is_temp: _cleanup_files(word_path)
        return Response({"error": "Please upload a DOCX/DOC file"}, status=400)

    _, pdf_path = _build_unique_paths(original_name, ".pdf")

    try:
        _docx_to_pdf_pure_python(word_path, pdf_path)
    except Exception as exc:
        _cleanup_files(pdf_path)
        return Response(
            {"error": f"Word to PDF conversion failed: {str(exc)}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(word_path)

    return Response({
        "message": "Word converted to PDF document",
        "file": settings.MEDIA_URL + pdf_path.name
    })


# ---------------- 3. PDF → PPT ----------------
@api_view(['POST'])
def pdf_to_ppt(request):
    pdf_path, original_name, is_temp = _get_single_input_path(request, 'file')

    if not pdf_path or not pdf_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if pdf_path.suffix.lower() != ".pdf":
        if is_temp: _cleanup_files(pdf_path)
        return Response({"error": "Please upload a PDF file"}, status=400)

    _, pptx_path = _build_unique_paths(original_name, ".pptx")

    try:
        pdf_doc = fitz.open(pdf_path)
        prs = Presentation()
        blank_slide_layout = prs.slide_layouts[6]

        if pdf_doc.page_count > 0:
            from pptx.util import Inches
            first_page = pdf_doc[0]
            prs.slide_width = Inches(first_page.rect.width / 72.0)
            prs.slide_height = Inches(first_page.rect.height / 72.0)

        for page in pdf_doc:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            img_bytes = pixmap.tobytes("png")
            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(
                io.BytesIO(img_bytes),
                left=0,
                top=0,
                width=prs.slide_width,
                height=prs.slide_height,
            )
            del pixmap, img_bytes

        pdf_doc.close()
        prs.save(str(pptx_path))
    except Exception as exc:
        _cleanup_files(pptx_path)
        return Response(
            {"error": f"PDF to PPT conversion failed: {str(exc)}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(pdf_path)

    return Response({
        "message": "PDF converted to PowerPoint presentation (.pptx)",
        "file": settings.MEDIA_URL + pptx_path.name
    })


# ---------------- 4. PPT → PDF ----------------
@api_view(['POST'])
def ppt_to_pdf(request):
    ppt_path, original_name, is_temp = _get_single_input_path(request, 'file')

    if not ppt_path or not ppt_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if ppt_path.suffix.lower() not in [".pptx", ".ppt"]:
        if is_temp: _cleanup_files(ppt_path)
        return Response({"error": "Please upload a PPT/PPTX file"}, status=400)

    _, pdf_path = _build_unique_paths(original_name, ".pdf")

    try:
        _pptx_to_pdf_pure_python(ppt_path, pdf_path)
    except Exception as exc:
        _cleanup_files(pdf_path)
        return Response(
            {"error": f"PPT to PDF conversion failed: {str(exc)}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(ppt_path)

    return Response({
        "message": "PowerPoint converted to PDF document",
        "file": settings.MEDIA_URL + pdf_path.name
    })


# ---------------- 5. IMAGE → PDF ----------------
@api_view(['POST'])
def image_to_pdf(request):
    input_items = _get_multiple_input_paths(request, 'files')

    if not input_items:
        return Response({"error": "No image files uploaded"}, status=400)

    out_name = f"images_converted-{uuid4().hex[:8]}.pdf"
    pdf_path = Path(settings.MEDIA_ROOT) / out_name
    pil_images = []

    try:
        for img_path, orig_name, is_temp in input_items:
            suffix = img_path.suffix.lower()
            if suffix not in [".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff"]:
                continue

            img = Image.open(img_path)
            if img.mode != "RGB":
                img = img.convert("RGB")
            pil_images.append(img)

        if not pil_images:
            return Response({"error": "No valid image files processed"}, status=400)

        first_img = pil_images[0]
        other_imgs = pil_images[1:] if len(pil_images) > 1 else []

        first_img.save(str(pdf_path), "PDF", resolution=100.0, save_all=True, append_images=other_imgs)

        for img in pil_images:
            img.close()

    except Exception as exc:
        _cleanup_files(pdf_path)
        return Response(
            {"error": f"Image to PDF conversion failed: {exc}"},
            status=500,
        )
    finally:
        for img_path, orig_name, is_temp in input_items:
            if is_temp:
                _cleanup_files(img_path)

    return Response({
        "message": f"Successfully converted {len(pil_images)} image(s) to PDF",
        "file": settings.MEDIA_URL + pdf_path.name
    })


# ---------------- 6. PDF → IMAGE ----------------
@api_view(['POST'])
def pdf_to_image(request):
    pdf_path, original_name, is_temp = _get_single_input_path(request, 'file')

    if not pdf_path or not pdf_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if pdf_path.suffix.lower() != ".pdf":
        if is_temp: _cleanup_files(pdf_path)
        return Response({"error": "Please upload a PDF file"}, status=400)

    try:
        pdf_doc = fitz.open(pdf_path)
        total_pages = pdf_doc.page_count

        if total_pages == 0:
            pdf_doc.close()
            raise ValueError("The uploaded PDF has no pages.")

        if total_pages == 1:
            png_name = f"{pdf_path.stem}-page1.png"
            png_path = Path(settings.MEDIA_ROOT) / png_name
            pixmap = pdf_doc[0].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            pixmap.save(str(png_path))
            pdf_doc.close()

            return Response({
                "message": "PDF page exported as high-res PNG image",
                "file": settings.MEDIA_URL + png_name
            })
        else:
            zip_name = f"{pdf_path.stem}-images.zip"
            zip_path = Path(settings.MEDIA_ROOT) / zip_name

            def render_page(idx):
                page = pdf_doc[idx]
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                return idx, pixmap.tobytes("png")

            max_workers = min(os.cpu_count() or 4, 8)
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                rendered_pages = list(executor.map(render_page, range(total_pages)))

            with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_STORED) as zipf:
                for idx, img_bytes in rendered_pages:
                    zipf.writestr(f"page-{idx+1}.png", img_bytes)

            pdf_doc.close()

            return Response({
                "message": f"Exported {total_pages} PDF pages into PNG Zip package",
                "file": settings.MEDIA_URL + zip_name
            })

    except Exception as exc:
        return Response(
            {"error": f"PDF to Image conversion failed: {exc}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(pdf_path)


# ---------------- 7. MERGE PDF ----------------
@api_view(['POST'])
def merge_pdf(request):
    input_items = _get_multiple_input_paths(request, 'files')

    if len(input_items) < 2:
        return Response({"error": "Please upload at least 2 PDF files to merge."}, status=400)

    out_name = f"merged-{uuid4().hex[:8]}.pdf"
    merged_pdf_path = Path(settings.MEDIA_ROOT) / out_name

    try:
        merged_doc = fitz.open()

        for pdf_path, orig_name, is_temp in input_items:
            if pdf_path.suffix.lower() != ".pdf":
                continue

            doc = fitz.open(pdf_path)
            merged_doc.insert_pdf(doc)
            doc.close()

        if merged_doc.page_count == 0:
            merged_doc.close()
            return Response({"error": "No valid PDF pages found to merge."}, status=400)

        merged_doc.save(str(merged_pdf_path))
        merged_doc.close()

    except Exception as exc:
        _cleanup_files(merged_pdf_path)
        return Response(
            {"error": f"Merge PDF failed: {exc}"},
            status=500,
        )
    finally:
        for pdf_path, orig_name, is_temp in input_items:
            if is_temp:
                _cleanup_files(pdf_path)

    return Response({
        "message": f"Successfully merged {len(input_items)} PDF files into one document",
        "file": settings.MEDIA_URL + merged_pdf_path.name
    })


# ---------------- 8. COMPRESS PDF ----------------
@api_view(['POST'])
def compress_pdf(request):
    pdf_path, original_name, is_temp = _get_single_input_path(request, 'file')
    mode = request.data.get("mode") or request.POST.get("mode", "extreme")

    if not pdf_path or not pdf_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if pdf_path.suffix.lower() != ".pdf":
        if is_temp: _cleanup_files(pdf_path)
        return Response({"error": "Please upload a PDF file"}, status=400)

    stem = Path(original_name).stem or "file"
    out_name = f"{stem}-compressed-{uuid4().hex[:8]}.pdf"
    compressed_path = Path(settings.MEDIA_ROOT) / out_name

    try:
        with open(pdf_path, "rb") as f:
            file_bytes = f.read()
        original_size = len(file_bytes)

        if original_size == 0:
            return Response({"error": "Uploaded PDF file is empty."}, status=400)

        pdf_doc = fitz.open(stream=file_bytes, filetype="pdf")

        if pdf_doc.page_count == 0:
            pdf_doc.close()
            return Response({"error": "PDF document contains no pages."}, status=400)

        is_extreme = mode == "extreme"
        dpi_val = 90 if is_extreme else 115
        quality_val = 26 if is_extreme else 48

        new_doc = fitz.open()

        def process_page(page_num):
            page = pdf_doc[page_num]
            pixmap = page.get_pixmap(dpi=dpi_val, alpha=False)
            img = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
            
            img_buffer = io.BytesIO()
            img.save(img_buffer, format="JPEG", quality=quality_val)
            img.close()
            return (page.rect.width, page.rect.height, img_buffer.getvalue())

        max_workers = min(os.cpu_count() or 4, 8)
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            results = list(executor.map(process_page, range(pdf_doc.page_count)))

        for w, h, img_bytes in results:
            new_page = new_doc.new_page(width=w, height=h)
            new_page.insert_image(new_page.rect, stream=img_bytes)

        compressed_bytes = new_doc.tobytes(garbage=4, deflate=True)
        new_doc.close()
        pdf_doc.close()

        compressed_size = len(compressed_bytes)
        if compressed_size >= original_size:
            compressed_bytes = file_bytes
            compressed_size = original_size
            savings_percent = 0.0
        else:
            savings_percent = round((1 - (compressed_size / max(original_size, 1))) * 100, 1)

        with open(compressed_path, "wb") as f:
            f.write(compressed_bytes)

    except Exception as exc:
        _cleanup_files(compressed_path)
        return Response(
            {"error": f"PDF compression failed: {str(exc)}"},
            status=400,
        )
    finally:
        if is_temp:
            _cleanup_files(pdf_path)

    return Response({
        "message": f"PDF compressed! Reduced by {savings_percent}%" if savings_percent > 0 else "PDF file is already fully optimized!",
        "file": settings.MEDIA_URL + compressed_path.name,
        "original_size": original_size,
        "compressed_size": compressed_size,
        "savings_percent": savings_percent,
    })


# ---------------- 9. COMPRESS IMAGE ----------------
@api_view(['POST'])
def compress_image(request):
    img_path, original_name, is_temp = _get_single_input_path(request, 'file')
    mode = request.data.get("mode") or request.POST.get("mode", "extreme")

    if not img_path or not img_path.exists():
        return Response({"error": "No image file uploaded"}, status=400)

    suffix = img_path.suffix.lower()
    if suffix not in [".jpg", ".jpeg", ".png", ".webp", ".bmp"]:
        if is_temp: _cleanup_files(img_path)
        return Response({"error": "Please upload a valid image file (JPG, PNG, WEBP, BMP)"}, status=400)

    stem = Path(original_name).stem or "image"
    out_name = f"{stem}-compressed-{uuid4().hex[:8]}{'.jpg' if mode == 'extreme' else suffix}"
    compressed_path = Path(settings.MEDIA_ROOT) / out_name

    try:
        with open(img_path, "rb") as f:
            raw_bytes = f.read()
        original_size = len(raw_bytes)

        img = Image.open(io.BytesIO(raw_bytes))
        is_extreme = mode == "extreme"

        quality = 28 if is_extreme else 50
        max_dim = 1000 if is_extreme else 1600

        img.thumbnail((max_dim, max_dim), Image.Resampling.BILINEAR)

        if img.mode in ("RGBA", "P") or is_extreme:
            img = img.convert("RGB")

        out_buffer = io.BytesIO()

        if is_extreme or suffix in [".jpg", ".jpeg", ".bmp"]:
            img.save(out_buffer, format="JPEG", quality=quality)
        elif suffix == ".webp":
            img.save(out_buffer, format="WEBP", quality=quality, method=4)
        else:
            img.save(out_buffer, format="PNG", compress_level=4)

        img.close()
        compressed_bytes = out_buffer.getvalue()

        if len(compressed_bytes) >= original_size:
            compressed_bytes = raw_bytes
            savings_percent = 0.0
        else:
            savings_percent = round((1 - (len(compressed_bytes) / max(original_size, 1))) * 100, 1)

        with open(compressed_path, "wb") as f:
            f.write(compressed_bytes)

        compressed_size = len(compressed_bytes)

    except Exception as exc:
        _cleanup_files(compressed_path)
        return Response(
            {"error": f"Image compression failed: {str(exc)}"},
            status=400,
        )
    finally:
        if is_temp:
            _cleanup_files(img_path)

    return Response({
        "message": f"Image compressed! Reduced by {savings_percent}%" if savings_percent > 0 else "Image is already fully optimized!",
        "file": settings.MEDIA_URL + compressed_path.name,
        "original_size": original_size,
        "compressed_size": compressed_size,
        "savings_percent": savings_percent,
    })


# ---------------- 10. COMPRESS WORD ----------------
@api_view(['POST'])
def compress_word(request):
    docx_path, original_name, is_temp = _get_single_input_path(request, 'file')
    mode = request.data.get("mode") or request.POST.get("mode", "extreme")

    if not docx_path or not docx_path.exists():
        return Response({"error": "No Word file uploaded"}, status=400)

    if docx_path.suffix.lower() not in [".docx", ".doc"]:
        if is_temp: _cleanup_files(docx_path)
        return Response({"error": "Please upload a Word file (.docx or .doc)"}, status=400)

    _, compressed_path = _build_unique_paths(original_name, docx_path.suffix.lower())

    try:
        original_size = docx_path.stat().st_size

        _compress_zip_media(docx_path, compressed_path, media_prefix="word/media/", extreme=(mode == "extreme"))

        compressed_size = compressed_path.stat().st_size
        if compressed_size >= original_size or compressed_size == 0:
            compressed_path.unlink(missing_ok=True)
            import shutil
            shutil.copy2(docx_path, compressed_path)
            compressed_size = original_size
            savings_percent = 0.0
        else:
            savings_percent = round((1 - (compressed_size / max(original_size, 1))) * 100, 1)

    except Exception as exc:
        _cleanup_files(compressed_path)
        return Response(
            {"error": f"Word compression failed: {str(exc)}"},
            status=400,
        )
    finally:
        if is_temp:
            _cleanup_files(docx_path)

    return Response({
        "message": f"Word document compressed! Size reduced by {savings_percent}%" if savings_percent > 0 else "Word document is already fully optimized!",
        "file": settings.MEDIA_URL + compressed_path.name,
        "original_size": original_size,
        "compressed_size": compressed_size,
        "savings_percent": savings_percent,
    })


# ---------------- 11. COMPRESS POWERPOINT ----------------
@api_view(['POST'])
def compress_ppt(request):
    pptx_path, original_name, is_temp = _get_single_input_path(request, 'file')
    mode = request.data.get("mode") or request.POST.get("mode", "extreme")

    if not pptx_path or not pptx_path.exists():
        return Response({"error": "No PowerPoint file uploaded"}, status=400)

    if pptx_path.suffix.lower() not in [".pptx", ".ppt"]:
        if is_temp: _cleanup_files(pptx_path)
        return Response({"error": "Please upload a PowerPoint file (.pptx or .ppt)"}, status=400)

    _, compressed_path = _build_unique_paths(original_name, pptx_path.suffix.lower())

    try:
        original_size = pptx_path.stat().st_size

        _compress_zip_media(pptx_path, compressed_path, media_prefix="ppt/media/", extreme=(mode == "extreme"))

        compressed_size = compressed_path.stat().st_size
        if compressed_size >= original_size or compressed_size == 0:
            compressed_path.unlink(missing_ok=True)
            import shutil
            shutil.copy2(pptx_path, compressed_path)
            compressed_size = original_size
            savings_percent = 0.0
        else:
            savings_percent = round((1 - (compressed_size / max(original_size, 1))) * 100, 1)

    except Exception as exc:
        _cleanup_files(compressed_path)
        return Response(
            {"error": f"PowerPoint compression failed: {str(exc)}"},
            status=400,
        )
    finally:
        if is_temp:
            _cleanup_files(pptx_path)

    return Response({
        "message": f"PowerPoint compressed! Size reduced by {savings_percent}%" if savings_percent > 0 else "PowerPoint deck is already fully optimized!",
        "file": settings.MEDIA_URL + compressed_path.name,
        "original_size": original_size,
        "compressed_size": compressed_size,
        "savings_percent": savings_percent,
    })


# ---------------- 12. EXCEL → PDF ----------------
@api_view(['POST'])
def excel_to_pdf(request):
    excel_path, original_name, is_temp = _get_single_input_path(request, 'file')

    if not excel_path or not excel_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if excel_path.suffix.lower() not in [".xlsx", ".xls"]:
        if is_temp: _cleanup_files(excel_path)
        return Response({"error": "Please upload an Excel file (.xlsx)"}, status=400)

    _, pdf_path = _build_unique_paths(original_name, ".pdf")

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        pdf_doc = fitz.open()

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            page = pdf_doc.new_page(width=792, height=612)
            
            lines = [f"Sheet: {sheet_name}", "=" * 50]
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(val) if val is not None else "" for val in row]
                if any(row_vals):
                    lines.append("  |  ".join(row_vals[:8]))

            full_text = "\n".join(lines[:35])
            page.insert_text(fitz.Point(36, 40), full_text, fontsize=11)

        wb.close()
        pdf_doc.save(str(pdf_path))
        pdf_doc.close()

    except Exception as exc:
        _cleanup_files(pdf_path)
        return Response(
            {"error": f"Excel to PDF conversion failed: {str(exc)}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(excel_path)

    return Response({
        "message": "Excel spreadsheet converted to PDF document successfully",
        "file": settings.MEDIA_URL + pdf_path.name
    })


# ---------------- 13. PDF → EXCEL ----------------
@api_view(['POST'])
def pdf_to_excel(request):
    pdf_path, original_name, is_temp = _get_single_input_path(request, 'file')

    if not pdf_path or not pdf_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if pdf_path.suffix.lower() != ".pdf":
        if is_temp: _cleanup_files(pdf_path)
        return Response({"error": "Please upload a PDF file"}, status=400)

    _, excel_path = _build_unique_paths(original_name, ".xlsx")

    try:
        pdf_doc = fitz.open(pdf_path)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

        for idx, page in enumerate(pdf_doc):
            ws = wb.create_sheet(title=f"Page {idx+1}")
            text_blocks = page.get_text("blocks")

            ws.append(["Block ID", "Extracted Content"])
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill

            for block_id, block in enumerate(text_blocks):
                content = block[4].strip()
                if content:
                    ws.append([f"Block {block_id+1}", content])

        pdf_doc.close()
        wb.save(str(excel_path))
        wb.close()

    except Exception as exc:
        _cleanup_files(excel_path)
        return Response(
            {"error": f"PDF to Excel conversion failed: {str(exc)}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(pdf_path)

    return Response({
        "message": "PDF content extracted into Excel workbook (.xlsx)",
        "file": settings.MEDIA_URL + excel_path.name
    })


# ---------------- 14. SPLIT PDF ----------------
@api_view(['POST'])
def split_pdf(request):
    pdf_path, original_name, is_temp = _get_single_input_path(request, 'file')

    if not pdf_path or not pdf_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if pdf_path.suffix.lower() != ".pdf":
        if is_temp: _cleanup_files(pdf_path)
        return Response({"error": "Please upload a PDF file"}, status=400)

    try:
        pdf_doc = fitz.open(pdf_path)
        total_pages = pdf_doc.page_count

        if total_pages == 0:
            pdf_doc.close()
            return Response({"error": "PDF has no pages to split."}, status=400)

        zip_name = f"{pdf_path.stem}-split.zip"
        zip_path = Path(settings.MEDIA_ROOT) / zip_name

        def split_page(idx):
            single_doc = fitz.open()
            single_doc.insert_pdf(pdf_doc, from_page=idx, to_page=idx)
            data = single_doc.tobytes(garbage=4, deflate=True)
            single_doc.close()
            return idx, data

        max_workers = min(os.cpu_count() or 4, 8)
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            rendered_pdfs = list(executor.map(split_page, range(total_pages)))

        with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zipf:
            for idx, pdf_bytes in rendered_pdfs:
                zipf.writestr(f"page-{idx+1}.pdf", pdf_bytes)

        pdf_doc.close()

    except Exception as exc:
        return Response(
            {"error": f"PDF split failed: {str(exc)}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(pdf_path)

    return Response({
        "message": f"Successfully split {total_pages} PDF pages into ZIP archive",
        "file": settings.MEDIA_URL + zip_name
    })


# ---------------- 15. ROTATE PDF ----------------
@api_view(['POST'])
def rotate_pdf(request):
    pdf_path, original_name, is_temp = _get_single_input_path(request, 'file')
    raw_angle = request.data.get("angle") or request.POST.get("angle", 90)
    try:
        angle = int(raw_angle)
    except (ValueError, TypeError):
        angle = 90

    if not pdf_path or not pdf_path.exists():
        return Response({"error": "No file uploaded"}, status=400)

    if pdf_path.suffix.lower() != ".pdf":
        if is_temp: _cleanup_files(pdf_path)
        return Response({"error": "Please upload a PDF file"}, status=400)

    _, rotated_path = _build_unique_paths(original_name, ".pdf")

    try:
        pdf_doc = fitz.open(pdf_path)

        for page in pdf_doc:
            page.set_rotation((page.rotation + angle) % 360)

        pdf_doc.save(str(rotated_path), deflate=True, garbage=3)
        pdf_doc.close()

    except Exception as exc:
        _cleanup_files(rotated_path)
        return Response(
            {"error": f"PDF rotation failed: {str(exc)}"},
            status=500,
        )
    finally:
        if is_temp:
            _cleanup_files(pdf_path)

    return Response({
        "message": f"PDF pages rotated by {angle}° successfully",
        "file": settings.MEDIA_URL + rotated_path.name
    })
